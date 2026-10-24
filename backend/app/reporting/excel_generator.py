"""
Automated Multi-Tab Excel Workbook Generator.
Uses openpyxl to produce a financial model and procurement audit workbook.
"""

from pathlib import Path
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..analytics.spot_buy_engine import analyze_spot_buys
from ..analytics.cycle_times import analyze_cycle_times
from ..analytics.three_way_match import analyze_three_way_match
from ..analytics.vendor_scorecard import analyze_vendors


def generate_p2p_excel_report(df: pd.DataFrame, output_path: str | Path | None = None) -> str:
    """
    Builds a styled multi-sheet Excel financial model workbook.
    """
    if output_path is None:
        base_dir = Path(__file__).resolve().parent.parent.parent
        reports_dir = base_dir / "data" / "reports"
        reports_dir.mkdir(parents=True, exist_ok=True)
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = reports_dir / f"P2P_Financial_Audit_Workbook_{timestamp_str}.xlsx"
    else:
        output_path = Path(output_path)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    accent_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # -------------------------------------------------------------
    # Sheet 1: Executive KPI Summary
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Executive Summary")
    ws1.views.sheetView[0].showGridLines = True
    ws1["A1"] = "Procure-to-Pay (P2P) Executive KPI & Compliance Summary"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Report Date: {datetime.now().strftime('%d %B %Y')} | Clean Data ETL Integration"
    ws1["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    spot_res = analyze_spot_buys(df)
    cycle_res = analyze_cycle_times(df)
    match_res = analyze_three_way_match(df)
    kpis = spot_res["kpis"]
    match_kpis = match_res["kpis"]

    kpi_rows = [
        ("Total Spend Audited (INR)", kpis["total_spend"], "Total gross spend across all processed purchase orders"),
        ("Total Purchase Orders", kpis["total_pos"], "Total number of completed PO cycles"),
        ("Contract Spend (INR)", kpis["contract_spend"], "Spend managed under active negotiated Master Service Agreements"),
        ("Contract Spend Rate (%)", f"{kpis['contract_spend_pct']}%", "Percentage of spend under negotiated enterprise contracts"),
        ("Spot Buy (Maverick) Spend (INR)", kpis["spot_spend"], "Uncontracted, ad-hoc, or one-off supplier purchases"),
        ("Spot Buy Spend Rate (%)", f"{kpis['spot_spend_pct']}%", "Percentage of total procurement spend occurring off-contract"),
        ("Spend Leakage / Premium (INR)", kpis["total_spend_leakage"], "Total cost penalty paid due to lack of negotiated contract tiers"),
        ("3-Way Match Clean Rate (%)", f"{match_kpis['matched_rate_pct']}%", "Percentage of invoices perfectly matching PO and Goods Receipt"),
        ("3-Way Discrepancy Exposure (INR)", match_kpis["total_disputed_exposure"], "Total value under active invoice hold or price dispute"),
        ("Avg Requisition-to-Approval (Days)", cycle_res["summary"]["cycle_time_creation_to_approval"]["mean_days"], "Managerial sign-off latency"),
        ("Avg Total P2P Cycle Time (Days)", cycle_res["summary"]["total_p2p_cycle_time"]["mean_days"], "Total elapsed duration from PO creation to payment"),
    ]

    ws1.cell(row=4, column=1, value="Procurement Metric").fill = header_fill
    ws1.cell(row=4, column=1).font = header_font
    ws1.cell(row=4, column=2, value="Value").fill = header_fill
    ws1.cell(row=4, column=2).font = header_font
    ws1.cell(row=4, column=3, value="Operational Definition / Context").fill = header_fill
    ws1.cell(row=4, column=3).font = header_font

    for i, (m, v, desc) in enumerate(kpi_rows, start=5):
        c1 = ws1.cell(row=i, column=1, value=m)
        c2 = ws1.cell(row=i, column=2, value=v)
        c3 = ws1.cell(row=i, column=3, value=desc)
        c1.font = bold_font
        c2.font = normal_font
        c3.font = normal_font
        if isinstance(v, (int, float)) and v > 100:
            c2.number_format = '#,##0.00'
        for c in (c1, c2, c3):
            c.border = thin_border

    # -------------------------------------------------------------
    # Sheet 2: Spot Buy & Maverick Governance
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Spot Buy Governance")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "Spot Buy vs. Contract Rule Breakdown & Maverick Spend"
    ws2["A1"].font = title_font

    rule_headers = ["Rule Trigger Category", "PO Count", "% of POs", "Spend Amount (INR)", "% of Total Spend"]
    for c_idx, h in enumerate(rule_headers, start=1):
        cell = ws2.cell(row=3, column=c_idx, value=h)
        cell.fill = accent_fill
        cell.font = header_font

    for r_idx, r in enumerate(spot_res["rule_breakdown"], start=4):
        ws2.cell(row=r_idx, column=1, value=r["rule"]).border = thin_border
        ws2.cell(row=r_idx, column=2, value=r["po_count"]).border = thin_border
        ws2.cell(row=r_idx, column=3, value=f"{r['po_pct']}%").border = thin_border
        c_spd = ws2.cell(row=r_idx, column=4, value=r["spend_amount"])
        c_spd.number_format = '#,##0.00'
        c_spd.border = thin_border
        ws2.cell(row=r_idx, column=5, value=f"{r['spend_pct']}%").border = thin_border

    # Top Spot Buy Suppliers Table
    ws2.cell(row=10, column=1, value="Top Spot Buy Suppliers by Spend Volume").font = Font(size=12, bold=True)
    vendor_headers = ["Vendor Name", "Vendor City", "Primary Category", "Total Spend (INR)", "PO Count", "Spend Leakage (INR)"]
    for c_idx, h in enumerate(vendor_headers, start=1):
        cell = ws2.cell(row=11, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for r_idx, v in enumerate(spot_res["top_spot_vendors"], start=12):
        ws2.cell(row=r_idx, column=1, value=v["vendor_name"]).border = thin_border
        ws2.cell(row=r_idx, column=2, value=v["vendor_city"]).border = thin_border
        ws2.cell(row=r_idx, column=3, value=v["primary_category"]).border = thin_border
        c_s = ws2.cell(row=r_idx, column=4, value=v["total_spend"])
        c_s.number_format = '#,##0.00'
        c_s.border = thin_border
        ws2.cell(row=r_idx, column=5, value=v["po_count"]).border = thin_border
        c_l = ws2.cell(row=r_idx, column=6, value=v["leakage"])
        c_l.number_format = '#,##0.00'
        c_l.border = thin_border

    # -------------------------------------------------------------
    # Sheet 3: 3-Way Match Exception Queue
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="3-Way Match Exceptions")
    ws3.views.sheetView[0].showGridLines = True
    ws3["A1"] = "3-Way Match Discrepancy Exception Queue"
    ws3["A1"].font = title_font

    match_headers = [
        "PO Number", "Invoice Number", "GRN Number", "Vendor Name", "Item Description",
        "PO Qty", "GRN Qty", "PO Price", "Invoice Price", "PO Total", "Invoice Total",
        "Disputed Value", "Discrepancy Status", "Recommended Resolution Action"
    ]
    for c_idx, h in enumerate(match_headers, start=1):
        cell = ws3.cell(row=3, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for r_idx, exc in enumerate(match_res["exception_queue"], start=4):
        ws3.cell(row=r_idx, column=1, value=exc["po_number"]).border = thin_border
        ws3.cell(row=r_idx, column=2, value=exc["invoice_number"]).border = thin_border
        ws3.cell(row=r_idx, column=3, value=exc["grn_number"]).border = thin_border
        ws3.cell(row=r_idx, column=4, value=exc["vendor_name"]).border = thin_border
        ws3.cell(row=r_idx, column=5, value=exc["item_name"]).border = thin_border
        ws3.cell(row=r_idx, column=6, value=exc["po_qty"]).border = thin_border
        ws3.cell(row=r_idx, column=7, value=exc["grn_qty"]).border = thin_border
        ws3.cell(row=r_idx, column=8, value=exc["po_unit_price"]).border = thin_border
        ws3.cell(row=r_idx, column=9, value=exc["invoice_unit_price"]).border = thin_border
        c_p = ws3.cell(row=r_idx, column=10, value=exc["po_total"])
        c_p.number_format = '#,##0.00'
        c_p.border = thin_border
        c_i = ws3.cell(row=r_idx, column=11, value=exc["invoice_total"])
        c_i.number_format = '#,##0.00'
        c_i.border = thin_border
        c_d = ws3.cell(row=r_idx, column=12, value=exc["disputed_value"])
        c_d.number_format = '#,##0.00'
        c_d.border = thin_border
        ws3.cell(row=r_idx, column=13, value=exc["status"]).border = thin_border
        ws3.cell(row=r_idx, column=14, value=exc["recommended_action"]).border = thin_border

    # -------------------------------------------------------------
    # Sheet 4: Supplier Performance & Kraljic Matrix
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Supplier Performance")
    ws4.views.sheetView[0].showGridLines = True
    ws4["A1"] = "Supplier OTIF Scorecard & Kraljic Matrix Segmentation"
    ws4["A1"].font = title_font

    vendor_res = analyze_vendors(df)
    supp_headers = [
        "Supplier Name", "City", "Category", "Kraljic Quadrant", "Risk Tier",
        "Performance Score", "Total Spend (INR)", "PO Count", "OTIF Rate (%)", "Avg Lead Time (Days)"
    ]
    for c_idx, h in enumerate(supp_headers, start=1):
        cell = ws4.cell(row=3, column=c_idx, value=h)
        cell.fill = accent_fill
        cell.font = header_font

    for r_idx, sc in enumerate(vendor_res["scorecard_table"], start=4):
        ws4.cell(row=r_idx, column=1, value=sc["vendor_name"]).border = thin_border
        ws4.cell(row=r_idx, column=2, value=sc["vendor_city"]).border = thin_border
        ws4.cell(row=r_idx, column=3, value=sc["primary_category"]).border = thin_border
        ws4.cell(row=r_idx, column=4, value=sc["kraljic_quadrant"]).border = thin_border
        ws4.cell(row=r_idx, column=5, value=sc["risk_tier"]).border = thin_border
        ws4.cell(row=r_idx, column=6, value=sc["performance_score"]).border = thin_border
        c_sp = ws4.cell(row=r_idx, column=7, value=sc["total_spend"])
        c_sp.number_format = '#,##0.00'
        c_sp.border = thin_border
        ws4.cell(row=r_idx, column=8, value=sc["po_count"]).border = thin_border
        ws4.cell(row=r_idx, column=9, value=f"{sc['otif_rate']}%").border = thin_border
        ws4.cell(row=r_idx, column=10, value=sc["avg_cycle_days"]).border = thin_border

    # Auto-fit column widths across all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

    wb.save(output_path)
    print(f"[Reporting Engine] Multi-tab Excel audit workbook generated at {output_path}")
    return str(output_path)


if __name__ == "__main__":
    from ..etl.pipeline import get_p2p_dataframe
    df = get_p2p_dataframe()
    generate_p2p_excel_report(df)
